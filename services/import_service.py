import csv
import io
import re
from werkzeug.datastructures import FileStorage
from models import db, User, UserRole, StudentProfile, FacultyProfile, Department, CollegeDepartment
from services.audit_service import log_audit_action

EMAIL_REGEX = re.compile(r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$')

def parse_file_rows(file_obj):
    """
    Parses an uploaded file (CSV or Excel) into a list of dictionaries with normalized header keys.
    """
    filename = file_obj.filename.lower()
    rows = []

    if filename.endswith('.csv'):
        stream = io.TextIOWrapper(file_obj.stream, encoding='utf-8-sig', errors='replace')
        reader = csv.DictReader(stream)
        for r in reader:
            normalized_row = {str(k).strip().lower().replace(' ', '_').replace('-', '_'): str(v).strip() for k, v in r.items() if k}
            if any(normalized_row.values()):
                rows.append(normalized_row)
    elif filename.endswith(('.xlsx', '.xls')):
        import openpyxl
        wb = openpyxl.load_workbook(file_obj.stream, data_only=True)
        sheet = wb.active
        headers = []
        for col_idx, cell in enumerate(sheet.iter_rows(min_row=1, max_row=1, values_only=True).__next__(), start=1):
            if cell:
                headers.append(str(cell).strip().lower().replace(' ', '_').replace('-', '_'))
            else:
                headers.append(f'col_{col_idx}')

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = {}
            for h, val in zip(headers, row):
                row_dict[h] = str(val).strip() if val is not None else ''
            if any(row_dict.values()):
                rows.append(row_dict)
    else:
        raise ValueError("Unsupported file format. Please upload a .csv, .xlsx, or .xls file.")

    return rows


def validate_import_data(rows, role):
    """
    Validates rows against role requirements, checks for missing fields, invalid formats,
    internal duplicates, and database duplicates.
    """
    role = role.upper()
    valid_rows = []
    skipped_rows = []
    errors = []

    seen_emails = set()
    seen_ids = set()

    for idx, row in enumerate(rows, start=2): # 1-indexed row number in spreadsheet
        row_errors = []

        # Common fields
        name = row.get('name') or row.get('full_name') or row.get('student_name') or row.get('faculty_name')
        email = row.get('email') or row.get('email_address')
        phone = row.get('phone') or row.get('mobile') or row.get('contact') or ''
        department = row.get('department') or row.get('dept') or row.get('branch')

        if not name:
            row_errors.append("Missing Name")
        if not email:
            row_errors.append("Missing Email")
        elif not EMAIL_REGEX.match(email):
            row_errors.append(f"Invalid email format '{email}'")

        if email:
            email_lower = email.lower()
            if email_lower in seen_emails:
                row_errors.append(f"Duplicate email in file '{email}'")
            seen_emails.add(email_lower)

            # Check existing in DB
            if User.query.filter_by(email=email_lower).first():
                row_errors.append(f"Email already registered in system: '{email}'")

        if not department:
            row_errors.append("Missing Department")
        else:
            dept_normalized = Department.normalize_code(department)
            if not Department.is_valid(dept_normalized):
                # Check DB departments
                if not CollegeDepartment.query.filter_by(code=dept_normalized).first():
                    row_errors.append(f"Invalid department '{department}'")
            department = dept_normalized

        parsed_data = {
            'row_num': idx,
            'name': name,
            'email': email.lower() if email else '',
            'phone': phone,
            'department': department,
            'raw': row
        }

        if role == UserRole.STUDENT:
            roll_number = row.get('roll_number') or row.get('roll_no') or row.get('usn') or row.get('reg_no')
            year_val = row.get('year') or row.get('academic_year') or '1'
            section = row.get('section') or row.get('sec') or 'A'

            if not roll_number:
                row_errors.append("Missing Roll Number / USN")
            else:
                roll_upper = roll_number.upper()
                if roll_upper in seen_ids:
                    row_errors.append(f"Duplicate roll number in file '{roll_number}'")
                seen_ids.add(roll_upper)

                if StudentProfile.query.filter_by(roll_number=roll_upper).first():
                    row_errors.append(f"Roll Number already registered: '{roll_number}'")
                parsed_data['roll_number'] = roll_upper

            try:
                year_int = int(float(year_val))
                if year_int not in (1, 2, 3, 4):
                    year_int = 1
            except Exception:
                year_int = 1

            parsed_data['year'] = year_int
            parsed_data['section'] = str(section).upper()[:10]

        elif role in (UserRole.FACULTY, UserRole.HOD, UserRole.FACULTY_ADMIN):
            emp_id = row.get('employee_id') or row.get('faculty_id') or row.get('emp_id')
            designation = row.get('designation') or row.get('role') or ('Head of Department' if role == UserRole.HOD else 'Assistant Professor')

            if not emp_id:
                row_errors.append("Missing Employee / Faculty ID")
            else:
                emp_upper = emp_id.upper()
                if emp_upper in seen_ids:
                    row_errors.append(f"Duplicate Employee ID in file '{emp_id}'")
                seen_ids.add(emp_upper)

                if FacultyProfile.query.filter_by(employee_id=emp_upper).first():
                    row_errors.append(f"Employee ID already registered: '{emp_id}'")
                parsed_data['employee_id'] = emp_upper

            parsed_data['designation'] = designation

        if row_errors:
            error_msg = f"Row {idx}: " + ", ".join(row_errors)
            errors.append(error_msg)
            parsed_data['errors'] = row_errors
            skipped_rows.append(parsed_data)
        else:
            valid_rows.append(parsed_data)

    return {
        'total_rows': len(rows),
        'valid_count': len(valid_rows),
        'skipped_count': len(skipped_rows),
        'valid_rows': valid_rows,
        'skipped_rows': skipped_rows,
        'errors': errors
    }


def execute_import(valid_rows, role, default_password='Pass@123', admin_user=None):
    """
    Inserts validated rows into the database atomically.
    """
    role = role.upper()
    created_count = 0
    created_users = []

    try:
        for row in valid_rows:
            user = User(
                name=row['name'],
                email=row['email'],
                phone=row.get('phone'),
                role=role,
                is_active=True
            )
            user.set_password(default_password)
            db.session.add(user)
            db.session.flush()

            if role == UserRole.STUDENT:
                profile = StudentProfile(
                    user_id=user.id,
                    roll_number=row['roll_number'],
                    department=row['department'],
                    year=row['year'],
                    section=row['section']
                )
                db.session.add(profile)
            elif role in (UserRole.FACULTY, UserRole.HOD, UserRole.FACULTY_ADMIN):
                profile = FacultyProfile(
                    user_id=user.id,
                    employee_id=row['employee_id'],
                    department=row['department'],
                    designation=row['designation']
                )
                db.session.add(profile)

                # If HOD, also assign to CollegeDepartment if department matches and department has no HOD
                if role == UserRole.HOD:
                    dept = CollegeDepartment.query.filter_by(code=row['department']).first()
                    if dept and not dept.hod_id:
                        dept.hod_id = user.id

            created_count += 1
            created_users.append(user.email)

        db.session.commit()

        if admin_user:
            log_audit_action(
                admin=admin_user,
                action="BULK_IMPORT",
                target_type="User",
                target_name=f"{created_count} {role} Users",
                details=f"Successfully imported {created_count} users of role {role}. Sample emails: {', '.join(created_users[:5])}"
            )

        return created_count, None
    except Exception as e:
        db.session.rollback()
        return 0, str(e)


def generate_sample_csv(role):
    """
    Generates a sample CSV template as a string for downloading.
    """
    role = role.upper()
    output = io.StringIO()
    writer = csv.writer(output)

    if role == UserRole.STUDENT:
        writer.writerow(['name', 'email', 'roll_number', 'department', 'year', 'section', 'phone'])
        writer.writerow(['Rahul Sharma', 'rahul.sharma@college.edu', '1MS24CS001', 'CSE', '1', 'A', '+91 9876543210'])
        writer.writerow(['Ananya Rao', 'ananya.rao@college.edu', '1MS24DS002', 'CSD', '1', 'B', '+91 9876543211'])
        writer.writerow(['Vikram Aditya', 'vikram.aditya@college.edu', '1MS24AI003', 'CSM', '1', 'A', '+91 9876543212'])
    elif role == UserRole.HOD:
        writer.writerow(['name', 'email', 'employee_id', 'department', 'designation', 'phone'])
        writer.writerow(['Dr. K. Ramanathan', 'hod.cse@college.edu', 'FAC-CSE-001', 'CSE', 'Head of Department & Professor', '+91 9840112233'])
        writer.writerow(['Dr. Ananya Roy', 'hod.csd@college.edu', 'FAC-CSD-002', 'CSD', 'Head of Department & Professor', '+91 9840112244'])
    else: # FACULTY
        writer.writerow(['name', 'email', 'employee_id', 'department', 'designation', 'phone'])
        writer.writerow(['Prof. Meenakshi', 'meenakshi@college.edu', 'FAC-IT-102', 'IT', 'Associate Professor', '+91 9840112255'])
        writer.writerow(['Dr. Arvind S.', 'arvind@college.edu', 'FAC-CSM-104', 'CSM', 'Assistant Professor', '+91 9840112266'])

    return output.getvalue()
